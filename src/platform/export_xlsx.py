from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from src.config import settings
from src.query.registry import registry

# 前十名拼在一个单元格里。列宽按最长的一个名字留，每人一行，避免默认窄列看起来像被截断。
_HOLDER_KEYS = frozenset({"holders", "top_holders"})
_HOLDER_WRAP = Alignment(wrap_text=True, vertical="top")
# 股本是股、市值是元。表上和 Excel 都改成亿，保留两位小数。
# 样例数据本身已经是亿（几十到几千），满 1 万才按原始单位换算，避免把样例再除一次。
_YI_KEYS = frozenset({"zgb", "ltgb", "mcap", "fcap"})
_YI = 100_000_000


def as_yi(value):
    if value is None or value == "":
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return value
    if number != number or number in (float("inf"), float("-inf")):
        return value
    scaled = number / _YI if abs(number) >= 10_000 else number
    return round(scaled, 2)


def _safe_token(name: str) -> str:
    cleaned = "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in (name or "pool"))
    return cleaned.strip("_") or "pool"


def _holder_lines(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return text
    parts = [part.strip() for part in text.split("、") if part.strip()]
    if len(parts) <= 1:
        return text
    return "\n".join(parts)


def _fit_holder_columns(ws, cols: list[int]) -> None:
    if not cols:
        return
    for col in cols:
        widest = 12
        for row_idx in range(2, ws.max_row + 1):
            cell = ws.cell(row_idx, col)
            cell.alignment = _HOLDER_WRAP
            text = "" if cell.value is None else str(cell.value)
            for line in text.split("\n"):
                widest = max(widest, len(line))
        ws.column_dimensions[get_column_letter(col)].width = min(widest * 2 + 4, 80)
    for row_idx in range(2, ws.max_row + 1):
        lines = 1
        for col in cols:
            text = ws.cell(row_idx, col).value
            if text:
                lines = max(lines, str(text).count("\n") + 1)
        if lines > 1:
            ws.row_dimensions[row_idx].height = min(16 * lines + 6, 220)


def write_query_xlsx(
    rows: list[dict],
    field_keys: list[str],
    pool_name: str,
    file_label: str = "",
) -> Path:
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    # 文件名 = 导出时间 + 用户确认的范围说明（如 沪市全部_看盘）。没传说明时仍用池名。
    filename = f"{stamp}_{_safe_token(file_label or pool_name)}.xlsx"
    path = settings.artifact_dir / filename
    wb = Workbook()
    ws = wb.active
    ws.title = "query"
    headers = [registry.get(k).label for k in field_keys]
    ws.append(headers)
    holder_cols = [i for i, key in enumerate(field_keys, start=1) if key in _HOLDER_KEYS]
    yi_cols = [i for i, key in enumerate(field_keys, start=1) if key in _YI_KEYS]
    for row in rows:
        values = []
        for key in field_keys:
            value = row.get(key)
            if key in _HOLDER_KEYS:
                value = _holder_lines(value)
            elif key in _YI_KEYS:
                value = as_yi(value)
            values.append(value)
        ws.append(values)
    for col in yi_cols:
        for row_idx in range(2, ws.max_row + 1):
            cell = ws.cell(row_idx, col)
            if isinstance(cell.value, float):
                cell.number_format = "0.00"
    _fit_holder_columns(ws, holder_cols)
    meta = wb.create_sheet("meta")
    meta.append(["exported_at", datetime.now().isoformat(timespec="seconds")])
    meta.append(["pool", pool_name])
    if file_label:
        meta.append(["export_label", file_label])
    meta.append(["fields", ",".join(field_keys)])
    wb.save(path)
    return path


def write_table_xlsx(headers: list, rows: list[dict], pool_name: str = "pasted") -> Path:
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    filename = f"{stamp}_{_safe_token(pool_name)}.xlsx"
    path = settings.artifact_dir / filename
    wb = Workbook()
    ws = wb.active
    ws.title = "query"
    labels = [str(h) for h in (headers or [])]
    if not labels and rows:
        labels = [str(k) for k in rows[0].keys()]
    ws.append(labels)
    for row in rows:
        ws.append([row.get(h) for h in labels])
    meta = wb.create_sheet("meta")
    meta.append(["exported_at", datetime.now().isoformat(timespec="seconds")])
    meta.append(["pool", pool_name])
    meta.append(["source", "pasted"])
    wb.save(path)
    return path


def write_agent_xlsx(question: str, answer: str, cites: list, tools: list) -> Path:
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    path = settings.artifact_dir / f"{stamp}_agent.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "agent"
    ws.append(["question", question])
    ws.append(["answer", answer])
    ws.append([])
    ws.append(["cites"])
    for cite in cites or []:
        if isinstance(cite, dict):
            ws.append([cite.get("cite") or cite.get("source"), cite.get("source")])
        else:
            ws.append([str(cite)])
    ws.append([])
    ws.append(["tools"])
    for tool in tools or []:
        if isinstance(tool, dict):
            ws.append([tool.get("id"), tool.get("ok"), tool.get("cite")])
        else:
            ws.append([str(tool)])
    wb.save(path)
    return path
