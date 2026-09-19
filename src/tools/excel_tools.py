from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from src.models import Artifact
from src.platform.export_xlsx import write_agent_xlsx
from src.platform.storage import record_artifact
from src.tools.base import ToolContext, ToolResult, ToolSpec
from src.tools.registry import registry


def _user_arts(ctx: ToolContext) -> list[Artifact]:
    return (
        ctx.db.query(Artifact)
        .filter_by(user_id=ctx.user_id)
        .order_by(Artifact.id.desc())
        .all()
    )


def _payload(item: Artifact) -> dict:
    return {
        "id": item.id,
        "filename": item.filename,
        "pool_name": item.pool_name,
        "path": item.path,
        "created_at": item.created_at.isoformat() if item.created_at else None,
        "exists": Path(item.path).exists(),
    }


def excel_list(args: dict, ctx: ToolContext) -> ToolResult:
    limit = int(args.get("limit") or 8)
    items = [_payload(x) for x in _user_arts(ctx)[:limit]]
    return ToolResult(ok=True, data=items, source="excel_list", cite="Excel 档案 · excel_list")


def _read_sheet(path: Path, max_rows: int = 20) -> dict:
    wb = load_workbook(path, data_only=True)
    ws = wb["query"] if "query" in wb.sheetnames else wb.active
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        rows.append({str(headers[i] or f"c{i}"): row[i] for i in range(len(headers))})
        if len(rows) >= max_rows:
            break
    meta = {}
    if "meta" in wb.sheetnames:
        for r in wb["meta"].iter_rows(values_only=True):
            if r and r[0]:
                meta[str(r[0])] = r[1]
    return {"headers": headers, "rows": rows, "meta": meta, "truncated": ws.max_row - 1 > len(rows)}


def excel_read(args: dict, ctx: ToolContext) -> ToolResult:
    items = _user_arts(ctx)
    if not items:
        return ToolResult(ok=False, error="还没有导出档案，请先在传统查询里导出 Excel", source="excel_read", cite="Excel 档案")
    target = None
    if args.get("id"):
        target = next((x for x in items if x.id == int(args["id"])), None)
    elif args.get("filename"):
        needle = str(args["filename"])
        target = next((x for x in items if needle in x.filename), None)
    else:
        target = items[0]
    if not target:
        return ToolResult(ok=False, error="找不到这份导出", source="excel_read", cite="Excel 档案")
    if not Path(target.path).exists():
        return ToolResult(ok=False, error=f"文件已不在磁盘: {target.filename}", source="excel_read", cite=target.filename)
    data = _read_sheet(Path(target.path), int(args.get("limit") or 20))
    data.update(_payload(target))
    return ToolResult(ok=True, data=data, source="excel_read", cite=f"Excel · {target.filename}")


def _index_rows(sheet: dict, key_label: str = "代码") -> dict[str, dict]:
    out = {}
    headers = [str(h or "") for h in sheet.get("headers") or []]
    key = key_label if key_label in headers else (headers[1] if len(headers) > 1 else (headers[0] if headers else ""))
    name_key = "名称" if "名称" in headers else ""
    for row in sheet.get("rows") or []:
        code = str(row.get(key) or "").split(".")[0]
        if code:
            out[code] = {"name": row.get(name_key, ""), "row": row}
    return out


def excel_diff(args: dict, ctx: ToolContext) -> ToolResult:
    items = _user_arts(ctx)
    if len(items) < 2:
        return ToolResult(ok=False, error="至少需要两份导出才能对照。请先再导出一次。", source="excel_diff", cite="Excel 对照")
    if args.get("id_a") and args.get("id_b"):
        a = next((x for x in items if x.id == int(args["id_a"])), None)
        b = next((x for x in items if x.id == int(args["id_b"])), None)
    else:
        a, b = items[1], items[0]
    if not a or not b:
        return ToolResult(ok=False, error="指定的档案不存在", source="excel_diff", cite="Excel 对照")
    col = args.get("column") or "前十大流通股东"
    left = _read_sheet(Path(a.path), 200)
    right = _read_sheet(Path(b.path), 200)
    idx_a, idx_b = _index_rows(left), _index_rows(right)
    changes = []
    for code in sorted(set(idx_a) | set(idx_b)):
        va = (idx_a.get(code) or {}).get("row", {}).get(col)
        vb = (idx_b.get(code) or {}).get("row", {}).get(col)
        if va != vb:
            changes.append(
                {
                    "code": code,
                    "name": (idx_b.get(code) or idx_a.get(code) or {}).get("name"),
                    "before": va,
                    "after": vb,
                }
            )
    return ToolResult(
        ok=True,
        data={
            "a": _payload(a),
            "b": _payload(b),
            "column": col,
            "changes": changes[:40],
            "change_count": len(changes),
        },
        source="excel_diff",
        cite=f"Excel 对照 · {a.filename} → {b.filename}",
    )


def excel_export(args: dict, ctx: ToolContext) -> ToolResult:
    path = write_agent_xlsx(
        question=args.get("question") or "",
        answer=args.get("answer") or "",
        cites=args.get("cites") or [],
        tools=args.get("tools") or [],
    )
    rec = record_artifact(
        ctx.db,
        user_id=ctx.user_id,
        path=path,
        pool_name="agent",
        field_keys=[],
        codes=[],
    )
    return ToolResult(ok=True, data=_payload(rec), source="excel_export", cite=f"Excel · {path.name}")


registry.register(ToolSpec("excel_list", "导出档案列表", "artifact", "列出用户最近导出的 Excel", {"type": "object", "properties": {"limit": {"type": "integer"}}}), excel_list)
registry.register(ToolSpec("excel_read", "读导出档案", "artifact", "按 id 或最新一份读取查询 Excel", {"type": "object", "properties": {"id": {"type": "integer"}, "filename": {"type": "string"}, "limit": {"type": "integer"}}}), excel_read)
registry.register(ToolSpec("excel_diff", "对照两份导出", "artifact", "按代码对齐比较某一列的变化", {"type": "object", "properties": {"id_a": {"type": "integer"}, "id_b": {"type": "integer"}, "column": {"type": "string"}}}), excel_diff)
registry.register(ToolSpec("excel_export", "导出本轮问答", "artifact", "把本轮问答写成 Excel 档案", {"type": "object", "properties": {"question": {"type": "string"}, "answer": {"type": "string"}, "cites": {"type": "array"}, "tools": {"type": "array"}}}), excel_export)
