from __future__ import annotations

import csv
import re
from html.parser import HTMLParser
from io import BytesIO, StringIO
from typing import Any

from openpyxl import load_workbook

CODE_RE = re.compile(r"\b(\d{6})(?:\.(SH|SZ|BJ|sh|sz|bj))?\b")
MAX_ROWS = 80


class _TableParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.tables: list[list[list[str]]] = []
        self._rows: list[list[str]] = []
        self._row: list[str] = []
        self._buf: list[str] = []
        self._in_cell = False

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag in {"td", "th"}:
            self._in_cell = True
            self._buf = []
        elif tag == "tr":
            self._row = []

    def handle_endtag(self, tag: str) -> None:
        if tag in {"td", "th"} and self._in_cell:
            self._row.append(re.sub(r"\s+", " ", "".join(self._buf)).strip())
            self._in_cell = False
        elif tag == "tr" and self._row:
            self._rows.append(self._row)
            self._row = []
        elif tag == "table" and self._rows:
            self.tables.append(self._rows)
            self._rows = []

    def handle_data(self, data: str) -> None:
        if self._in_cell:
            self._buf.append(data)


def looks_like_table(text: str) -> bool:
    blob = (text or "").strip()
    if not blob:
        return False
    if "<table" in blob.lower():
        return True
    lines = [ln for ln in blob.splitlines() if ln.strip()]
    if len(lines) < 2:
        return False
    if any("\t" in ln for ln in lines[:3]):
        return True
    if sum(1 for ln in lines[:6] if "," in ln) >= 2:
        return True
    coded = sum(1 for ln in lines if CODE_RE.search(ln))
    return coded >= 2


def extract_codes(text: str) -> list[str]:
    seen: set[str] = set()
    out: list[str] = []
    for match in CODE_RE.finditer(text or ""):
        code = match.group(1)
        suffix = (match.group(2) or "").upper()
        token = f"{code}.{suffix}" if suffix else code
        if code in seen:
            continue
        seen.add(code)
        out.append(token)
    return out


def _matrix_to_sheet(matrix: list[list[str]], max_rows: int = MAX_ROWS) -> dict[str, Any]:
    if not matrix:
        return {"headers": [], "rows": [], "codes": [], "truncated": False}
    width = max(len(r) for r in matrix)
    padded = [r + [""] * (width - len(r)) for r in matrix]
    headers = [str(h or f"c{i}") for i, h in enumerate(padded[0])]
    body = padded[1:]
    truncated = len(body) > max_rows
    rows = []
    for raw in body[:max_rows]:
        rows.append({headers[i]: raw[i] for i in range(width)})
    blob = "\n".join("\t".join(r) for r in padded)
    return {
        "headers": headers,
        "rows": rows,
        "codes": extract_codes(blob),
        "truncated": truncated,
        "row_count": len(body),
    }


def parse_table_text(text: str, max_rows: int = MAX_ROWS) -> dict[str, Any]:
    blob = (text or "").strip()
    if not blob:
        return {"headers": [], "rows": [], "codes": [], "truncated": False, "row_count": 0}
    if "<table" in blob.lower():
        parser = _TableParser()
        parser.feed(blob)
        if parser.tables:
            return _matrix_to_sheet(parser.tables[0], max_rows)
    lines = [ln.rstrip("\r") for ln in blob.splitlines() if ln.strip()]
    if not lines:
        return {"headers": [], "rows": [], "codes": [], "truncated": False, "row_count": 0}
    if any("\t" in ln for ln in lines[:4]):
        matrix = [ln.split("\t") for ln in lines]
        return _matrix_to_sheet(matrix, max_rows)
    sample = "\n".join(lines[:8])
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;|")
        reader = csv.reader(StringIO("\n".join(lines)), dialect)
        return _matrix_to_sheet([list(r) for r in reader], max_rows)
    except csv.Error:
        matrix = [re.split(r"\s{2,}|\s", ln.strip()) for ln in lines]
        return _matrix_to_sheet(matrix, max_rows)


def parse_xlsx_bytes(raw: bytes, max_rows: int = MAX_ROWS) -> dict[str, Any]:
    wb = load_workbook(BytesIO(raw), data_only=True)
    ws = wb["query"] if "query" in wb.sheetnames else wb.active
    matrix: list[list[str]] = []
    for row in ws.iter_rows(values_only=True):
        matrix.append(["" if cell is None else str(cell) for cell in row])
        if len(matrix) > max_rows + 1:
            break
    sheet = _matrix_to_sheet(matrix, max_rows)
    sheet["truncated"] = (ws.max_row or 0) - 1 > len(sheet["rows"])
    return sheet
