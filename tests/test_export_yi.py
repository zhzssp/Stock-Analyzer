from openpyxl import load_workbook

from src.platform.export_xlsx import as_yi, write_query_xlsx


def test_as_yi_converts_raw_shares_and_yuan_but_keeps_sample_yi():
    assert as_yi(19_405_918_198) == 194.06
    assert as_yi(22_604_455_882) == 226.04
    assert as_yi(412.5) == 412.5
    assert as_yi(8.6) == 8.6
    assert as_yi(None) is None
    assert as_yi("") is None


def test_query_excel_writes_yi_with_two_decimals(monkeypatch, tmp_path):
    from src.config import settings

    monkeypatch.setattr(settings, "data_dir", tmp_path)
    settings.artifact_dir.mkdir(parents=True, exist_ok=True)
    path = write_query_xlsx(
        [{"name": "平安银行", "zgb": 19_405_918_198, "ltgb": 19_405_684_991, "mcap": 22_604_455_882, "fcap": 18_442_544_043}],
        ["name", "zgb", "ltgb", "mcap", "fcap"],
        "自选",
    )
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    assert [ws.cell(1, col).value for col in range(1, 6)] == ["名称", "总股本(亿)", "流通股本(亿)", "总市值(亿)", "流通市值(亿)"]
    assert ws.cell(2, 2).value == 194.06
    assert ws.cell(2, 3).value == 194.06
    assert ws.cell(2, 4).value == 226.04
    assert ws.cell(2, 5).value == 184.43
    assert ws.cell(2, 4).number_format == "0.00"
    wb.close()
